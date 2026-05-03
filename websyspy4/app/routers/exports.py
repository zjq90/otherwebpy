"""
数据导出路由
包含按时间、热点、产品、分类等导出数据的API接口
"""

from datetime import date, datetime
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from typing import Optional, List
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from app.database import get_db
from app.models.user import User
from app.models.product import Product
from app.models.category import Category
from app.models.reservation import Reservation, ReservationStatus
from app.routers.users import get_current_admin

router = APIRouter()


def create_excel_workbook(title: str, headers: List[str], rows: List[List]):
    """
    创建Excel工作簿
    
    Args:
        title: 工作表标题
        headers: 表头列表
        rows: 数据行列表
        
    Returns:
        BytesIO: Excel文件字节流
    """
    output = BytesIO()
    
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = title
    
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for col_idx, header in enumerate(headers, 1):
        cell = worksheet.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.fill = header_fill
        cell.border = thin_border
    
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, cell_value in enumerate(row_data, 1):
            cell = worksheet.cell(row=row_idx, column=col_idx, value=cell_value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for col_idx in range(1, len(headers) + 1):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18
    
    workbook.save(output)
    output.seek(0)
    
    return output


@router.get("/reservations")
def export_reservations(
    start_date: Optional[date] = Query(None, description="开始日期"),
    end_date: Optional[date] = Query(None, description="结束日期"),
    status: Optional[str] = Query(None, description="预约状态"),
    product_id: Optional[int] = Query(None, description="产品ID"),
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """
    导出预约数据（管理员权限）
    支持按时间范围、状态、产品过滤
    
    Args:
        start_date: 开始日期
        end_date: 结束日期
        status: 预约状态
        product_id: 产品ID
        current_admin: 当前管理员
        db: 数据库会话
        
    Returns:
        StreamingResponse: Excel文件下载
    """
    query = db.query(Reservation).filter(Reservation.is_deleted == False)
    
    if start_date:
        query = query.filter(func.date(Reservation.created_at) >= start_date)
    if end_date:
        query = query.filter(func.date(Reservation.created_at) <= end_date)
    if status:
        query = query.filter(Reservation.status == status)
    if product_id:
        query = query.filter(Reservation.product_id == product_id)
    
    reservations = query.order_by(Reservation.created_at.desc()).all()
    
    headers = ["预约单号", "用户手机号", "用户姓名", "产品名称", "开始日期", "结束日期", 
               "租借天数", "数量", "日租金", "总租金", "押金", "押金状态", "预约状态", "创建时间"]
    
    rows = []
    for r in reservations:
        row = [
            r.reservation_no,
            r.user.phone if r.user else "",
            r.user.name if r.user and r.user.name else "",
            r.product.name if r.product else "",
            str(r.start_date),
            str(r.end_date),
            r.rental_days,
            r.quantity,
            float(r.daily_rent),
            float(r.total_rent),
            float(r.deposit),
            "已支付" if r.deposit_paid else "未支付",
            get_status_text(r.status),
            r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else ""
        ]
        rows.append(row)
    
    output = create_excel_workbook("预约数据", headers, rows)
    
    filename = f"reservations_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/products")
def export_products(
    category_id: Optional[int] = Query(None, description="分类ID"),
    is_hot: Optional[bool] = Query(None, description="是否热门"),
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """
    导出产品数据（管理员权限）
    支持按分类、是否热门过滤
    
    Args:
        category_id: 分类ID
        is_hot: 是否热门
        current_admin: 当前管理员
        db: 数据库会话
        
    Returns:
        StreamingResponse: Excel文件下载
    """
    query = db.query(Product).filter(Product.is_deleted == False)
    
    if category_id:
        query = query.filter(Product.category_id == category_id)
    if is_hot is not None:
        query = query.filter(Product.is_hot == is_hot)
    
    products = query.order_by(Product.created_at.desc()).all()
    
    headers = ["产品名称", "分类名称", "日租金", "押金", "库存数量", "可用数量", "状态", "是否热门", "创建时间"]
    
    rows = []
    for p in products:
        row = [
            p.name,
            p.category.name if p.category else "",
            float(p.daily_rent),
            float(p.deposit),
            p.stock_quantity,
            p.available_quantity,
            "上架" if p.status == 1 else "下架",
            "是" if p.is_hot else "否",
            p.created_at.strftime("%Y-%m-%d %H:%M:%S") if p.created_at else ""
        ]
        rows.append(row)
    
    output = create_excel_workbook("产品数据", headers, rows)
    
    filename = f"products_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/categories")
def export_categories(
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """
    导出分类数据（管理员权限）
    
    Args:
        current_admin: 当前管理员
        db: 数据库会话
        
    Returns:
        StreamingResponse: Excel文件下载
    """
    categories = db.query(Category).filter(
        Category.is_deleted == False
    ).order_by(Category.sort_order.asc()).all()
    
    headers = ["分类名称", "分类描述", "排序", "产品数量", "创建时间"]
    
    rows = []
    for c in categories:
        product_count = db.query(Product).filter(
            Product.category_id == c.id,
            Product.is_deleted == False
        ).count()
        
        row = [
            c.name,
            c.description or "",
            c.sort_order,
            product_count,
            c.created_at.strftime("%Y-%m-%d %H:%M:%S") if c.created_at else ""
        ]
        rows.append(row)
    
    output = create_excel_workbook("分类数据", headers, rows)
    
    filename = f"categories_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/hot-products")
def export_hot_products(
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """
    导出热门产品数据（管理员权限）
    
    Args:
        current_admin: 当前管理员
        db: 数据库会话
        
    Returns:
        StreamingResponse: Excel文件下载
    """
    query = db.query(Product).filter(
        Product.is_deleted == False,
        Product.is_hot == True
    )
    
    products = query.order_by(Product.created_at.desc()).all()
    
    headers = ["产品名称", "分类名称", "日租金", "押金", "库存数量", "可用数量", "状态", "创建时间"]
    
    rows = []
    for p in products:
        row = [
            p.name,
            p.category.name if p.category else "",
            float(p.daily_rent),
            float(p.deposit),
            p.stock_quantity,
            p.available_quantity,
            "上架" if p.status == 1 else "下架",
            p.created_at.strftime("%Y-%m-%d %H:%M:%S") if p.created_at else ""
        ]
        rows.append(row)
    
    output = create_excel_workbook("热门产品数据", headers, rows)
    
    filename = f"hot_products_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/reservation-stats")
def export_reservation_stats(
    start_date: Optional[date] = Query(None, description="开始日期"),
    end_date: Optional[date] = Query(None, description="结束日期"),
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """
    导出预约统计数据（管理员权限）
    按状态统计预约数量和金额
    
    Args:
        start_date: 开始日期
        end_date: 结束日期
        current_admin: 当前管理员
        db: 数据库会话
        
    Returns:
        StreamingResponse: Excel文件下载
    """
    query = db.query(Reservation).filter(Reservation.is_deleted == False)
    
    if start_date:
        query = query.filter(func.date(Reservation.created_at) >= start_date)
    if end_date:
        query = query.filter(func.date(Reservation.created_at) <= end_date)
    
    reservations = query.all()
    
    status_stats = {}
    for r in reservations:
        status = get_status_text(r.status)
        if status not in status_stats:
            status_stats[status] = {"count": 0, "total_rent": 0, "total_deposit": 0}
        status_stats[status]["count"] += 1
        status_stats[status]["total_rent"] += float(r.total_rent)
        status_stats[status]["total_deposit"] += float(r.deposit)
    
    headers = ["预约状态", "预约数量", "租金总额", "押金总额"]
    
    rows = []
    for status, stats in status_stats.items():
        row = [
            status,
            stats["count"],
            round(stats["total_rent"], 2),
            round(stats["total_deposit"], 2)
        ]
        rows.append(row)
    
    output = create_excel_workbook("预约统计数据", headers, rows)
    
    filename = f"reservation_stats_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


def get_status_text(status: str) -> str:
    """
    获取预约状态的中文描述
    
    Args:
        status: 状态字符串
        
    Returns:
        str: 状态中文描述
    """
    status_map = {
        ReservationStatus.PENDING: "待确认",
        ReservationStatus.CONFIRMED: "已确认",
        ReservationStatus.PAID: "已支付押金",
        ReservationStatus.PICKED_UP: "已领取",
        ReservationStatus.RETURNED: "已归还",
        ReservationStatus.COMPLETED: "已完成",
        ReservationStatus.CANCELLED: "已取消"
    }
    return status_map.get(status, status)
