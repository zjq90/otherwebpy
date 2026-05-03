"""
FastAPI主应用入口文件
包含应用配置、中间件、路由注册等
"""

import os
from pathlib import Path
from datetime import date, datetime
from decimal import Decimal
from typing import Optional
from fastapi import FastAPI, Depends, HTTPException, status, Request, Form, UploadFile, File, Query
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

import models
import schemas
import crud
from database import engine, get_db, init_db
from config import (
    PRODUCT_IMAGES_DIR, EXCEL_EXPORT_DIR, 
    ALLOWED_IMAGE_TYPES, MAX_IMAGE_SIZE
)

# 初始化数据库表
init_db()

# 创建FastAPI应用实例
app = FastAPI(
    title="进销存管理系统",
    description="基于Python+FastAPI+SQLite+Bootstrap的进销存管理系统",
    version="1.0.0"
)

# 挂载静态文件目录
static_dir = Path(__file__).resolve().parent / "static"
static_dir.mkdir(exist_ok=True)
app.mount("/static", StaticFiles(directory=str(static_dir)), name="static")

# 挂载上传文件目录
app.mount("/uploads", StaticFiles(directory=str(PRODUCT_IMAGES_DIR.parent)), name="uploads")

# 配置模板目录
templates_dir = Path(__file__).resolve().parent / "templates"
templates_dir.mkdir(exist_ok=True)
templates = Jinja2Templates(directory=str(templates_dir))

# 确保导出目录存在
EXCEL_EXPORT_DIR.mkdir(parents=True, exist_ok=True)


# ==================== 页面路由 ====================

@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    """
    首页/登录页
    """
    return templates.TemplateResponse("login.html", {"request": request})


@app.get("/dashboard", response_class=HTMLResponse)
async def dashboard_page(request: Request, db: Session = Depends(get_db)):
    """
    数据看板页面
    """
    stats = crud.get_dashboard_stats(db)
    return templates.TemplateResponse("dashboard.html", {
        "request": request,
        "stats": stats
    })


@app.get("/users", response_class=HTMLResponse)
async def users_page(request: Request):
    """
    用户管理页面
    """
    return templates.TemplateResponse("users.html", {"request": request})


@app.get("/categories", response_class=HTMLResponse)
async def categories_page(request: Request):
    """
    分类管理页面
    """
    return templates.TemplateResponse("categories.html", {"request": request})


@app.get("/products", response_class=HTMLResponse)
async def products_page(request: Request):
    """
    产品管理页面
    """
    return templates.TemplateResponse("products.html", {"request": request})


@app.get("/inventory", response_class=HTMLResponse)
async def inventory_page(request: Request):
    """
    进销存管理页面
    """
    return templates.TemplateResponse("inventory.html", {"request": request})


# ==================== API路由 - 用户管理 ====================

@app.get("/api/users")
async def get_users(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    keyword: str = None,
    db: Session = Depends(get_db)
):
    """
    获取用户列表（分页）
    """
    users, total = crud.get_users(db, page=page, page_size=page_size, keyword=keyword)
    total_pages = (total + page_size - 1) // page_size
    
    return {
        "items": [schemas.UserResponse.model_validate(u) for u in users],
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages
    }


@app.get("/api/users/{user_id}")
async def get_user(user_id: int, db: Session = Depends(get_db)):
    """
    获取单个用户
    """
    user = crud.get_user(db, user_id)
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    return schemas.UserResponse.model_validate(user)


@app.post("/api/users")
async def create_user(user: schemas.UserCreate, db: Session = Depends(get_db)):
    """
    创建用户
    """
    existing_user = crud.get_user_by_username(db, user.username)
    if existing_user:
        raise HTTPException(status_code=400, detail="用户名已存在")
    
    db_user = crud.create_user(db, user)
    return {"success": True, "message": "创建成功", "data": {"id": db_user.id}}


@app.put("/api/users/{user_id}")
async def update_user(user_id: int, user: schemas.UserUpdate, db: Session = Depends(get_db)):
    """
    更新用户
    """
    db_user = crud.update_user(db, user_id, user)
    if not db_user:
        raise HTTPException(status_code=404, detail="用户不存在")
    return {"success": True, "message": "更新成功"}


@app.delete("/api/users/{user_id}")
async def delete_user(user_id: int, db: Session = Depends(get_db)):
    """
    删除用户
    """
    success = crud.delete_user(db, user_id)
    if not success:
        raise HTTPException(status_code=404, detail="用户不存在或无法删除")
    return {"success": True, "message": "删除成功"}


# ==================== API路由 - 分类管理 ====================

@app.get("/api/categories")
async def get_categories(
    parent_id: Optional[int] = None,
    db: Session = Depends(get_db)
):
    """
    获取分类列表
    """
    categories = crud.get_all_categories(db)
    return {"items": [schemas.CategoryResponse.model_validate(c) for c in categories]}


@app.get("/api/categories/{category_id}")
async def get_category(category_id: int, db: Session = Depends(get_db)):
    """
    获取单个分类
    """
    category = crud.get_category(db, category_id)
    if not category:
        raise HTTPException(status_code=404, detail="分类不存在")
    return schemas.CategoryResponse.model_validate(category)


@app.post("/api/categories")
async def create_category(category: schemas.CategoryCreate, db: Session = Depends(get_db)):
    """
    创建分类
    """
    db_category = crud.create_category(db, category)
    return {"success": True, "message": "创建成功", "data": {"id": db_category.id}}


@app.put("/api/categories/{category_id}")
async def update_category(
    category_id: int, 
    category: schemas.CategoryUpdate, 
    db: Session = Depends(get_db)
):
    """
    更新分类
    """
    db_category = crud.update_category(db, category_id, category)
    if not db_category:
        raise HTTPException(status_code=404, detail="分类不存在")
    return {"success": True, "message": "更新成功"}


@app.delete("/api/categories/{category_id}")
async def delete_category(category_id: int, db: Session = Depends(get_db)):
    """
    删除分类
    """
    success = crud.delete_category(db, category_id)
    if not success:
        raise HTTPException(
            status_code=400, 
            detail="分类不存在或存在子分类/关联产品，无法删除"
        )
    return {"success": True, "message": "删除成功"}


# ==================== API路由 - 产品管理 ====================

@app.get("/api/products")
async def get_products(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    keyword: str = None,
    category_id: int = None,
    status: int = None,
    db: Session = Depends(get_db)
):
    """
    获取产品列表（分页）
    """
    products, total = crud.get_products(
        db, page=page, page_size=page_size,
        keyword=keyword, category_id=category_id, status=status
    )
    total_pages = (total + page_size - 1) // page_size
    
    return {
        "items": [schemas.ProductResponse.model_validate(p) for p in products],
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages
    }


@app.get("/api/products/{product_id}")
async def get_product(product_id: int, db: Session = Depends(get_db)):
    """
    获取单个产品
    """
    product = crud.get_product(db, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="产品不存在")
    return schemas.ProductResponse.model_validate(product)


@app.post("/api/products")
async def create_product(product: schemas.ProductCreate, db: Session = Depends(get_db)):
    """
    创建产品
    """
    existing_product = crud.get_product_by_code(db, product.code)
    if existing_product:
        raise HTTPException(status_code=400, detail="产品编码已存在")
    
    db_product = crud.create_product(db, product)
    return {"success": True, "message": "创建成功", "data": {"id": db_product.id}}


@app.put("/api/products/{product_id}")
async def update_product(
    product_id: int, 
    product: schemas.ProductUpdate, 
    db: Session = Depends(get_db)
):
    """
    更新产品
    """
    existing_product = crud.get_product(db, product_id)
    if not existing_product:
        raise HTTPException(status_code=404, detail="产品不存在")
    
    if product.code is not None:
        product_by_code = crud.get_product_by_code(db, product.code)
        if product_by_code and product_by_code.id != product_id:
            raise HTTPException(status_code=400, detail="产品编码已存在")
    
    db_product = crud.update_product(db, product_id, product)
    if not db_product:
        raise HTTPException(status_code=404, detail="产品不存在")
    return {"success": True, "message": "更新成功"}


@app.put("/api/products/{product_id}/toggle-status")
async def toggle_product_status(product_id: int, db: Session = Depends(get_db)):
    """
    切换产品上下架状态
    """
    db_product = crud.toggle_product_status(db, product_id)
    if not db_product:
        raise HTTPException(status_code=404, detail="产品不存在")
    return {
        "success": True, 
        "message": "状态切换成功", 
        "data": {"status": db_product.status}
    }


@app.post("/api/products/{product_id}/upload-image")
async def upload_product_image(
    product_id: int,
    image: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """
    上传产品图片
    """
    product = crud.get_product(db, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="产品不存在")
    
    # 验证文件类型
    if image.content_type not in ALLOWED_IMAGE_TYPES:
        raise HTTPException(
            status_code=400,
            detail=f"不支持的文件类型，仅支持: {', '.join(ALLOWED_IMAGE_TYPES)}"
        )
    
    # 读取文件内容并验证大小
    image_data = await image.read()
    if len(image_data) > MAX_IMAGE_SIZE:
        raise HTTPException(
            status_code=400,
            detail=f"文件大小超过限制，最大允许: {MAX_IMAGE_SIZE // 1024 // 1024}MB"
        )
    
    # 生成文件名
    file_ext = image.filename.split('.')[-1] if image.filename else 'jpg'
    file_name = f"product_{product_id}_{datetime.now().strftime('%Y%m%d%H%M%S')}.{file_ext}"
    file_path = PRODUCT_IMAGES_DIR / file_name
    
    # 保存文件
    with open(file_path, "wb") as f:
        f.write(image_data)
    
    # 更新产品图片路径
    image_url = f"/uploads/products/{file_name}"
    crud.update_product_image(db, product_id, image_url)
    
    return {
        "success": True,
        "message": "上传成功",
        "data": {"image_url": image_url}
    }


@app.delete("/api/products/{product_id}")
async def delete_product(product_id: int, db: Session = Depends(get_db)):
    """
    删除产品
    """
    success = crud.delete_product(db, product_id)
    if not success:
        raise HTTPException(
            status_code=400, 
            detail="产品不存在或存在进销存记录，无法删除"
        )
    return {"success": True, "message": "删除成功"}


@app.get("/api/products/export/excel")
async def export_products_excel(
    keyword: str = None,
    category_id: int = None,
    db: Session = Depends(get_db)
):
    """
    导出产品列表为Excel
    """
    products = crud.get_products_for_export(db, keyword=keyword, category_id=category_id)
    
    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "产品列表"
    
    # 定义样式
    header_font = Font(bold=True, size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 表头
    headers = [
        "ID", "产品名称", "产品编码", "分类", "规格", "单位",
        "成本价", "销售价", "库存数量", "状态", "创建时间"
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # 数据行
    for row_idx, product in enumerate(products, 2):
        category_name = product.category.name if product.category else ""
        status_text = "上架" if product.status == 1 else "下架"
        
        row_data = [
            product.id,
            product.name,
            product.code,
            category_name,
            product.specification or "",
            product.unit or "",
            float(product.cost_price) if product.cost_price else 0,
            float(product.selling_price) if product.selling_price else 0,
            product.stock_quantity,
            status_text,
            product.created_at.strftime("%Y-%m-%d %H:%M:%S") if product.created_at else ""
        ]
        
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
    
    # 调整列宽
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    # 保存文件
    file_name = f"products_export_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    file_path = EXCEL_EXPORT_DIR / file_name
    wb.save(file_path)
    
    return FileResponse(
        path=file_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=file_name
    )


# ==================== API路由 - 进销存管理 ====================

@app.get("/api/inventory")
async def get_inventory_records(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    product_id: Optional[int] = Query(None),
    record_type: Optional[str] = Query(None),
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    """
    获取进销存记录列表（分页）
    """
    # 解析日期参数
    parsed_start_date = None
    parsed_end_date = None
    
    if start_date and start_date.strip():
        try:
            parsed_start_date = datetime.strptime(start_date, "%Y-%m-%d").date()
        except ValueError:
            pass
    
    if end_date and end_date.strip():
        try:
            parsed_end_date = datetime.strptime(end_date, "%Y-%m-%d").date()
        except ValueError:
            pass
    
    records, total = crud.get_inventory_records(
        db, page=page, page_size=page_size,
        product_id=product_id, record_type=record_type,
        start_date=parsed_start_date, end_date=parsed_end_date
    )
    total_pages = (total + page_size - 1) // page_size
    
    return {
        "items": [schemas.InventoryRecordResponse.model_validate(r) for r in records],
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages
    }


@app.post("/api/inventory")
async def create_inventory_record(
    record: schemas.InventoryRecordCreate,
    db: Session = Depends(get_db)
):
    """
    创建进销存记录（入库/出库）
    """
    db_record = crud.create_inventory_record(db, record)
    if not db_record:
        raise HTTPException(
            status_code=400,
            detail="创建失败，请检查产品是否存在或库存是否充足"
        )
    return {"success": True, "message": "操作成功", "data": {"id": db_record.id}}


# ==================== API路由 - 数据看板 ====================

@app.get("/api/dashboard/stats")
async def get_dashboard_stats(db: Session = Depends(get_db)):
    """
    获取数据看板统计数据
    """
    stats = crud.get_dashboard_stats(db)
    return stats


# ==================== 测试数据生成 ====================

@app.post("/api/test/generate-data")
async def generate_test_data(db: Session = Depends(get_db)):
    """
    生成测试数据
    """
    from datetime import timedelta
    import random
    
    # 生成管理员用户
    admin_user = crud.get_user_by_username(db, "admin")
    if not admin_user:
        admin_data = schemas.UserCreate(
            username="admin",
            password="admin123",
            real_name="管理员",
            email="admin@example.com",
            phone="13800138000",
            role="admin"
        )
        admin_user = crud.create_user(db, admin_data)
    
    # 生成普通用户
    test_users = [
        {"username": "user1", "real_name": "张三", "role": "user"},
        {"username": "user2", "real_name": "李四", "role": "user"},
        {"username": "user3", "real_name": "王五", "role": "user"},
    ]
    
    for user_data in test_users:
        existing = crud.get_user_by_username(db, user_data["username"])
        if not existing:
            user_create = schemas.UserCreate(
                username=user_data["username"],
                password="123456",
                real_name=user_data["real_name"],
                role=user_data["role"]
            )
            crud.create_user(db, user_create)
    
    # 生成分类
    categories_data = [
        {"name": "电子产品", "parent_id": None, "sort_order": 1},
        {"name": "办公用品", "parent_id": None, "sort_order": 2},
        {"name": "生活用品", "parent_id": None, "sort_order": 3},
        {"name": "手机", "parent_id": None, "sort_order": 4},
        {"name": "笔记本电脑", "parent_id": None, "sort_order": 5},
        {"name": "外设设备", "parent_id": None, "sort_order": 6},
    ]
    
    created_categories = []
    for cat_data in categories_data:
        cat_create = schemas.CategoryCreate(**cat_data)
        db_cat = crud.create_category(db, cat_create)
        created_categories.append(db_cat)
    
    # 生成产品
    products_data = [
        {"name": "iPhone 15 Pro", "code": "IP15P-001", "category_id": 4, 
         "specification": "256GB", "unit": "台", "cost_price": 7999, "selling_price": 8999, "stock_quantity": 50, "min_stock": 10},
        {"name": "MacBook Pro 14", "code": "MBP14-001", "category_id": 5,
         "specification": "M3 Pro/16GB/512GB", "unit": "台", "cost_price": 14999, "selling_price": 16999, "stock_quantity": 20, "min_stock": 5},
        {"name": "Logitech 无线鼠标", "code": "LOGI-M001", "category_id": 6,
         "specification": "MX Master 3S", "unit": "个", "cost_price": 599, "selling_price": 799, "stock_quantity": 100, "min_stock": 20},
        {"name": "华为Mate 60 Pro", "code": "HW-M60P-001", "category_id": 4,
         "specification": "256GB", "unit": "台", "cost_price": 6499, "selling_price": 6999, "stock_quantity": 30, "min_stock": 10},
        {"name": "戴尔显示器27寸", "code": "DELL-D27-001", "category_id": 1,
         "specification": "4K/IPS", "unit": "台", "cost_price": 2999, "selling_price": 3599, "stock_quantity": 15, "min_stock": 5},
        {"name": "A4打印纸", "code": "PAPER-A4-001", "category_id": 2,
         "specification": "70g/500张", "unit": "包", "cost_price": 25, "selling_price": 35, "stock_quantity": 500, "min_stock": 100},
        {"name": "晨光中性笔", "code": "PEN-CG-001", "category_id": 2,
         "specification": "0.5mm黑色", "unit": "支", "cost_price": 1.5, "selling_price": 2.5, "stock_quantity": 1000, "min_stock": 200},
        {"name": "小米手环8", "code": "MI-B8-001", "category_id": 1,
         "specification": "标准版", "unit": "个", "cost_price": 249, "selling_price": 299, "stock_quantity": 80, "min_stock": 20},
        {"name": "AirPods Pro 2", "code": "APP-AP2-001", "category_id": 6,
         "specification": "USB-C", "unit": "副", "cost_price": 1799, "selling_price": 1999, "stock_quantity": 40, "min_stock": 10},
        {"name": "键盘机械键盘", "code": "KB-MECH-001", "category_id": 6,
         "specification": "青轴/87键", "unit": "个", "cost_price": 399, "selling_price": 499, "stock_quantity": 60, "min_stock": 15},
    ]
    
    created_products = []
    for prod_data in products_data:
        existing = crud.get_product_by_code(db, prod_data["code"])
        if not existing:
            prod_create = schemas.ProductCreate(**prod_data)
            db_prod = crud.create_product(db, prod_create)
            created_products.append(db_prod)
        else:
            created_products.append(existing)
    
    # 生成进销存记录
    users = crud.get_users(db, page=1, page_size=100)[0]
    user_ids = [u.id for u in users]
    
    # 生成最近30天的记录
    for i in range(50):
        product = random.choice(created_products)
        record_type = random.choice(["in", "out"])
        quantity = random.randint(1, 20)
        unit_price = product.cost_price if record_type == "in" else product.selling_price
        
        record_data = schemas.InventoryRecordCreate(
            product_id=product.id,
            record_type=record_type,
            quantity=quantity,
            unit_price=unit_price,
            operator_id=random.choice(user_ids) if user_ids else None,
            remark=f"测试{'入库' if record_type == 'in' else '出库'}记录",
            reference_no=f"{'RK' if record_type == 'in' else 'CK'}{datetime.now().strftime('%Y%m%d')}{random.randint(1000, 9999)}"
        )
        
        # 手动更新库存和创建记录（避免库存不足问题）
        db_product = crud.get_product(db, product.id)
        if db_product:
            if record_type == "in":
                db_product.stock_quantity += quantity
            else:
                # 出库时确保库存足够
                if db_product.stock_quantity >= quantity:
                    db_product.stock_quantity -= quantity
                else:
                    continue
            
            total_amount = quantity * unit_price
            db_record = models.InventoryRecord(
                product_id=product.id,
                record_type=record_type,
                quantity=quantity,
                unit_price=unit_price,
                total_amount=total_amount,
                operator_id=record_data.operator_id,
                remark=record_data.remark,
                reference_no=record_data.reference_no,
                created_at=datetime.now() - timedelta(days=random.randint(0, 30))
            )
            db.add(db_record)
            db.commit()
    
    return {
        "success": True,
        "message": "测试数据生成成功！\n管理员账号: admin / admin123\n普通用户账号: user1 / 123456"
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
