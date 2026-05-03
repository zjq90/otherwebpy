"""
考勤管理API路由
提供考勤设置、考勤记录生成、查询和报表导出功能
"""
from io import BytesIO
from typing import Optional, List
from datetime import datetime, date, timedelta, time
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import and_, desc, func
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from database import get_db
from models.attendance_setting import AttendanceSetting
from models.attendance_record import AttendanceRecord
from models.access_record import AccessRecord
from models.staff import Staff

router = APIRouter(prefix="/api/attendance", tags=["考勤管理"])


def get_or_create_setting(db: Session) -> AttendanceSetting:
    """
    获取或创建考勤设置
    如果没有设置，创建默认设置
    """
    setting = db.query(AttendanceSetting).first()
    if not setting:
        setting = AttendanceSetting()
        db.add(setting)
        db.commit()
        db.refresh(setting)
    return setting


def parse_time_str(time_str: str) -> time:
    """
    解析时间字符串为time对象
    """
    return datetime.strptime(time_str, "%H:%M").time()


def calculate_attendance_status(
    check_time: Optional[datetime],
    expected_time_str: str,
    grace_minutes: int,
    is_check_in: bool = True
) -> int:
    """
    计算考勤状态
    返回: 0-缺勤, 1-正常, 2-迟到, 3-早退
    """
    if not check_time:
        return 0
    
    expected_time = parse_time_str(expected_time_str)
    check_time_only = check_time.time()
    
    # 计算时间差（分钟）
    expected_dt = datetime.combine(date.today(), expected_time)
    check_dt = datetime.combine(date.today(), check_time_only)
    diff_minutes = (check_dt - expected_dt).total_seconds() / 60
    
    if is_check_in:
        # 签到：晚于预期时间+宽容时间 = 迟到
        if diff_minutes > grace_minutes:
            return 2
        return 1
    else:
        # 签退：早于预期时间-宽容时间 = 早退
        if diff_minutes < -grace_minutes:
            return 3
        return 1


def calculate_day_status(morning_status: int, afternoon_status: int, enable_afternoon: bool) -> int:
    """
    计算当日总状态
    """
    if not enable_afternoon:
        return morning_status
    
    # 组合上午和下午状态
    if morning_status == 0 and afternoon_status == 0:
        return 0  # 全天缺勤
    elif morning_status == 2 and afternoon_status == 3:
        return 4  # 迟到+早退
    elif morning_status in [2, 4] or afternoon_status in [2, 4]:
        return 2  # 迟到
    elif morning_status == 3 or afternoon_status == 3:
        return 3  # 早退
    return 1  # 正常


@router.get("/setting")
def get_attendance_setting(db: Session = Depends(get_db)):
    """
    获取考勤设置
    """
    setting = get_or_create_setting(db)
    return {
        "code": 200,
        "message": "success",
        "data": setting.to_dict()
    }


@router.put("/setting")
def update_attendance_setting(
    morning_start_time: Optional[str] = None,
    morning_end_time: Optional[str] = None,
    afternoon_start_time: Optional[str] = None,
    afternoon_end_time: Optional[str] = None,
    grace_minutes: Optional[int] = None,
    enable_afternoon: Optional[int] = None,
    work_days: Optional[str] = None,
    remark: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    更新考勤设置
    """
    setting = get_or_create_setting(db)
    
    update_data = {
        "morning_start_time": morning_start_time,
        "morning_end_time": morning_end_time,
        "afternoon_start_time": afternoon_start_time,
        "afternoon_end_time": afternoon_end_time,
        "grace_minutes": grace_minutes,
        "enable_afternoon": enable_afternoon,
        "work_days": work_days,
        "remark": remark
    }
    
    for key, value in update_data.items():
        if value is not None:
            setattr(setting, key, value)
    
    db.commit()
    db.refresh(setting)
    
    return {
        "code": 200,
        "message": "更新成功",
        "data": setting.to_dict()
    }


@router.post("/generate")
def generate_attendance_records(
    start_date: str,
    end_date: str,
    db: Session = Depends(get_db)
):
    """
    生成考勤记录
    根据出入记录和考勤设置，为指定日期范围生成考勤记录
    """
    setting = get_or_create_setting(db)
    
    try:
        start_dt = datetime.strptime(start_date, "%Y-%m-%d").date()
        end_dt = datetime.strptime(end_date, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="日期格式错误，请使用YYYY-MM-DD格式")
    
    if start_dt > end_dt:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    
    # 获取所有在职员工
    staff_list = db.query(Staff).filter(Staff.status == 1).all()
    if not staff_list:
        return {
            "code": 200,
            "message": "没有需要处理的员工",
            "data": {"generated_count": 0}
        }
    
    generated_count = 0
    
    # 遍历每个日期
    current_date = start_dt
    while current_date <= end_dt:
        # 检查是否是工作日
        weekday = current_date.isoweekday()  # 1=周一, 7=周日
        work_days_list = [int(d) for d in setting.work_days.split(",") if d.strip().isdigit()]
        
        if weekday not in work_days_list:
            current_date += timedelta(days=1)
            continue
        
        # 为每个员工生成考勤记录
        for staff in staff_list:
            # 检查是否已存在该日期的考勤记录
            existing = db.query(AttendanceRecord).filter(
                AttendanceRecord.staff_id == staff.id,
                AttendanceRecord.attendance_date == current_date
            ).first()
            
            if existing:
                continue
            
            # 查询当天的出入记录
            day_start = datetime.combine(current_date, time.min)
            day_end = datetime.combine(current_date, time.max)
            
            # 上午时段的记录
            morning_end = datetime.combine(current_date, parse_time_str(setting.morning_end_time))
            morning_records = db.query(AccessRecord).filter(
                AccessRecord.staff_id == staff.id,
                AccessRecord.access_time >= day_start,
                AccessRecord.access_time <= morning_end,
                AccessRecord.verify_result == 1
            ).order_by(AccessRecord.access_time).all()
            
            # 下午时段的记录
            afternoon_start = datetime.combine(current_date, parse_time_str(setting.afternoon_start_time))
            afternoon_records = db.query(AccessRecord).filter(
                AccessRecord.staff_id == staff.id,
                AccessRecord.access_time > afternoon_start,
                AccessRecord.access_time <= day_end,
                AccessRecord.verify_result == 1
            ).order_by(AccessRecord.access_time).all()
            
            # 确定签到签退时间
            morning_check_in = None
            morning_check_out = None
            afternoon_check_in = None
            afternoon_check_out = None
            
            # 上午：最早的进入记录作为签到，最晚的离开记录作为签退
            for record in morning_records:
                if record.access_type == 1 and (morning_check_in is None or record.access_time < morning_check_in):
                    morning_check_in = record.access_time
                elif record.access_type == 0 and (morning_check_out is None or record.access_time > morning_check_out):
                    morning_check_out = record.access_time
            
            # 下午：最早的进入记录作为签到，最晚的离开记录作为签退
            for record in afternoon_records:
                if record.access_type == 1 and (afternoon_check_in is None or record.access_time < afternoon_check_in):
                    afternoon_check_in = record.access_time
                elif record.access_type == 0 and (afternoon_check_out is None or record.access_time > afternoon_check_out):
                    afternoon_check_out = record.access_time
            
            # 计算考勤状态
            morning_status = calculate_attendance_status(
                morning_check_in, setting.morning_start_time, setting.grace_minutes, is_check_in=True
            )
            # 如果没有签到，但有签退，也认为是异常
            if morning_status == 0 and morning_check_out:
                morning_status = 2  # 视为迟到（没有签到）
            
            afternoon_status = 0
            if setting.enable_afternoon:
                afternoon_status = calculate_attendance_status(
                    afternoon_check_out, setting.afternoon_end_time, setting.grace_minutes, is_check_in=False
                )
                if afternoon_status == 0 and afternoon_check_in:
                    afternoon_status = 3  # 视为早退（没有签退）
            
            # 计算当日总状态
            day_status = calculate_day_status(morning_status, afternoon_status, setting.enable_afternoon)
            
            # 计算工作时长
            work_duration = 0
            if morning_check_in and morning_check_out:
                work_duration += int((morning_check_out - morning_check_in).total_seconds() / 60)
            if afternoon_check_in and afternoon_check_out:
                work_duration += int((afternoon_check_out - afternoon_check_in).total_seconds() / 60)
            
            # 创建考勤记录
            attendance_record = AttendanceRecord(
                staff_id=staff.id,
                attendance_date=current_date,
                morning_check_in=morning_check_in,
                morning_check_out=morning_check_out,
                afternoon_check_in=afternoon_check_in,
                afternoon_check_out=afternoon_check_out,
                morning_status=morning_status,
                afternoon_status=afternoon_status,
                day_status=day_status,
                work_duration=work_duration
            )
            
            db.add(attendance_record)
            generated_count += 1
        
        current_date += timedelta(days=1)
    
    db.commit()
    
    return {
        "code": 200,
        "message": "生成成功",
        "data": {"generated_count": generated_count}
    }


@router.get("/list")
def get_attendance_record_list(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    staff_name: Optional[str] = None,
    staff_no: Optional[str] = None,
    department: Optional[str] = None,
    day_status: Optional[int] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    获取考勤记录列表
    支持分页、多条件筛选
    """
    query = db.query(AttendanceRecord)
    
    # 按人员姓名筛选
    if staff_name:
        query = query.join(Staff).filter(Staff.name.contains(staff_name))
    
    # 按工号筛选
    if staff_no:
        query = query.join(Staff).filter(Staff.staff_no.contains(staff_no))
    
    # 按部门筛选
    if department:
        query = query.join(Staff).filter(Staff.department == department)
    
    # 按状态筛选
    if day_status is not None:
        query = query.filter(AttendanceRecord.day_status == day_status)
    
    # 按日期范围筛选
    if start_date:
        try:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d").date()
            query = query.filter(AttendanceRecord.attendance_date >= start_dt)
        except ValueError:
            pass
    
    if end_date:
        try:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d").date()
            query = query.filter(AttendanceRecord.attendance_date <= end_dt)
        except ValueError:
            pass
    
    # 统计总数
    total = query.count()
    
    # 分页查询（按日期倒序）
    records = query.order_by(desc(AttendanceRecord.attendance_date)).offset((page - 1) * page_size).limit(page_size).all()
    
    return {
        "code": 200,
        "message": "success",
        "data": {
            "list": [record.to_dict() for record in records],
            "total": total,
            "page": page,
            "page_size": page_size
        }
    }


@router.get("/statistics")
def get_attendance_statistics(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    department: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    获取考勤统计数据
    """
    query = db.query(AttendanceRecord)
    
    # 应用筛选条件
    if department:
        query = query.join(Staff).filter(Staff.department == department)
    
    if start_date:
        try:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d").date()
            query = query.filter(AttendanceRecord.attendance_date >= start_dt)
        except ValueError:
            pass
    
    if end_date:
        try:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d").date()
            query = query.filter(AttendanceRecord.attendance_date <= end_dt)
        except ValueError:
            pass
    
    # 统计各状态数量
    total = query.count()
    normal_count = query.filter(AttendanceRecord.day_status == 1).count()
    late_count = query.filter(AttendanceRecord.day_status == 2).count()
    early_count = query.filter(AttendanceRecord.day_status == 3).count()
    late_early_count = query.filter(AttendanceRecord.day_status == 4).count()
    absent_count = query.filter(AttendanceRecord.day_status == 0).count()
    
    # 计算总工作时长
    total_duration = db.query(func.sum(AttendanceRecord.work_duration)).filter(
        AttendanceRecord.id.in_([r.id for r in query.all()])
    ).scalar() or 0
    
    return {
        "code": 200,
        "message": "success",
        "data": {
            "total": total,
            "normal_count": normal_count,
            "late_count": late_count,
            "early_count": early_count,
            "late_early_count": late_early_count,
            "absent_count": absent_count,
            "total_duration_minutes": total_duration,
            "total_duration_hours": round(total_duration / 60, 2)
        }
    }


@router.get("/export/report")
def export_attendance_report(
    report_type: str = Query(..., description="报表类型：month-月报, year-年报, custom-自定义时间段"),
    year: Optional[int] = None,
    month: Optional[int] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    department: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    导出考勤报表
    支持月报、年报、自定义时间段报表
    """
    query = db.query(AttendanceRecord)
    
    # 确定日期范围
    if report_type == "month":
        if not year or not month:
            raise HTTPException(status_code=400, detail="月报需要指定年份和月份")
        try:
            start_dt = date(year, month, 1)
            if month == 12:
                end_dt = date(year + 1, 1, 1) - timedelta(days=1)
            else:
                end_dt = date(year, month + 1, 1) - timedelta(days=1)
            query = query.filter(
                AttendanceRecord.attendance_date >= start_dt,
                AttendanceRecord.attendance_date <= end_dt
            )
        except ValueError:
            raise HTTPException(status_code=400, detail="日期参数错误")
    
    elif report_type == "year":
        if not year:
            raise HTTPException(status_code=400, detail="年报需要指定年份")
        try:
            start_dt = date(year, 1, 1)
            end_dt = date(year, 12, 31)
            query = query.filter(
                AttendanceRecord.attendance_date >= start_dt,
                AttendanceRecord.attendance_date <= end_dt
            )
        except ValueError:
            raise HTTPException(status_code=400, detail="年份参数错误")
    
    elif report_type == "custom":
        if not start_date or not end_date:
            raise HTTPException(status_code=400, detail="自定义报表需要指定开始日期和结束日期")
        try:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d").date()
            end_dt = datetime.strptime(end_date, "%Y-%m-%d").date()
            query = query.filter(
                AttendanceRecord.attendance_date >= start_dt,
                AttendanceRecord.attendance_date <= end_dt
            )
        except ValueError:
            raise HTTPException(status_code=400, detail="日期格式错误")
    
    else:
        raise HTTPException(status_code=400, detail="不支持的报表类型")
    
    # 按部门筛选
    if department:
        query = query.join(Staff).filter(Staff.department == department)
    
    records = query.order_by(AttendanceRecord.staff_id, AttendanceRecord.attendance_date).all()
    
    # 创建Excel工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "考勤报表"
    
    # 定义样式
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    cell_alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入表头
    headers = ["序号", "工号", "姓名", "部门", "考勤日期", "上午签到", "上午签退", "上午状态", 
               "下午签到", "下午签退", "下午状态", "当日状态", "工作时长"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # 写入数据
    for row_idx, record in enumerate(records, 2):
        record_dict = record.to_dict()
        row_data = [
            row_idx - 1,
            record_dict.get("staff_no", ""),
            record_dict.get("staff_name", ""),
            record_dict.get("department", ""),
            record_dict.get("attendance_date", ""),
            record_dict.get("morning_check_in", ""),
            record_dict.get("morning_check_out", ""),
            record_dict.get("morning_status_text", ""),
            record_dict.get("afternoon_check_in", ""),
            record_dict.get("afternoon_check_out", ""),
            record_dict.get("afternoon_status_text", ""),
            record_dict.get("day_status_text", ""),
            record_dict.get("work_duration_text", "")
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = cell_alignment
            cell.border = thin_border
    
    # 设置列宽
    column_widths = [8, 12, 12, 15, 12, 10, 10, 10, 10, 10, 10, 12, 12]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    # 保存到内存
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # 生成文件名
    if report_type == "month":
        filename = f"考勤月报_{year}年{month}月.xlsx"
    elif report_type == "year":
        filename = f"考勤年报_{year}年.xlsx"
    else:
        filename = f"考勤报表_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
