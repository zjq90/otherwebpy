"""
出入记录API路由
提供出入记录的查询、导出和人脸识别验证功能
"""
import os
import uuid
import random
from io import BytesIO
from typing import Optional
from datetime import datetime, date
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import and_, desc
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from database import get_db
from models.access_record import AccessRecord
from models.staff import Staff
from models.device import Device
from config import CAPTURE_PHOTO_DIR, FACE_SIMILARITY_THRESHOLD

router = APIRouter(prefix="/api/access", tags=["出入记录管理"])


@router.get("/list")
def get_access_record_list(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    staff_name: Optional[str] = None,
    staff_no: Optional[str] = None,
    device_id: Optional[int] = None,
    access_type: Optional[int] = None,
    verify_result: Optional[int] = None,
    start_time: Optional[str] = None,
    end_time: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    获取出入记录列表
    支持分页、多条件筛选
    """
    query = db.query(AccessRecord)
    
    # 按人员姓名筛选
    if staff_name:
        query = query.join(Staff).filter(Staff.name.contains(staff_name))
    
    # 按工号筛选
    if staff_no:
        query = query.join(Staff).filter(Staff.staff_no.contains(staff_no))
    
    # 按设备筛选
    if device_id:
        query = query.filter(AccessRecord.device_id == device_id)
    
    # 按出入类型筛选
    if access_type is not None:
        query = query.filter(AccessRecord.access_type == access_type)
    
    # 按验证结果筛选
    if verify_result is not None:
        query = query.filter(AccessRecord.verify_result == verify_result)
    
    # 按时间范围筛选
    if start_time:
        try:
            start_dt = datetime.strptime(start_time, "%Y-%m-%d")
            query = query.filter(AccessRecord.access_time >= start_dt)
        except ValueError:
            pass
    
    if end_time:
        try:
            end_dt = datetime.strptime(end_time, "%Y-%m-%d")
            end_dt = end_dt.replace(hour=23, minute=59, second=59)
            query = query.filter(AccessRecord.access_time <= end_dt)
        except ValueError:
            pass
    
    # 统计总数
    total = query.count()
    
    # 分页查询（按时间倒序），预加载关联对象
    records = query.options(
        joinedload(AccessRecord.staff),
        joinedload(AccessRecord.device)
    ).order_by(desc(AccessRecord.access_time)).offset((page - 1) * page_size).limit(page_size).all()
    
    return {
        "code": 200,
        "message": "success",
        "data": {
            "list": [record.to_dict() for record in records],
            "total": total,
            "page": page,
            "page_size": page_size
        }
    }


@router.get("/{record_id}")
def get_access_record_detail(record_id: int, db: Session = Depends(get_db)):
    """
    获取出入记录详情
    """
    record = db.query(AccessRecord).options(
        joinedload(AccessRecord.staff),
        joinedload(AccessRecord.device)
    ).filter(AccessRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="记录不存在")
    
    return {
        "code": 200,
        "message": "success",
        "data": record.to_dict()
    }


@router.post("/verify")
def face_verify(
    device_id: int,
    staff_id: Optional[int] = None,
    capture_photo: Optional[UploadFile] = File(None),
    db: Session = Depends(get_db)
):
    """
    人脸识别验证接口
    模拟门禁设备扫脸验证流程
    """
    # 检查设备是否存在且在线
    device = db.query(Device).filter(Device.id == device_id).first()
    if not device:
        raise HTTPException(status_code=404, detail="设备不存在")
    if device.status != 1:
        raise HTTPException(status_code=400, detail="设备不在线")
    
    # 如果提供了staff_id，直接验证（模拟已知人员）
    if staff_id:
        staff = db.query(Staff).filter(Staff.id == staff_id, Staff.status == 1).first()
        if not staff:
            # 验证失败，创建失败记录
            record = AccessRecord(
                staff_id=0,
                device_id=device_id,
                access_type=1,
                access_time=datetime.now(),
                similarity=0.0,
                verify_result=0,
                remark="人员不存在或已离职"
            )
            db.add(record)
            db.commit()
            return {
                "code": 400,
                "message": "验证失败：人员不存在或已离职",
                "data": {"verify_result": 0, "similarity": 0.0}
            }
        
        # 模拟人脸识别相似度（0.85-0.99之间）
        similarity = round(random.uniform(0.85, 0.99), 2)
        
        # 处理抓拍照片
        capture_photo_path = None
        if capture_photo:
            file_ext = os.path.splitext(capture_photo.filename)[1] or ".jpg"
            file_name = f"{uuid.uuid4()}{file_ext}"
            file_path = os.path.join(CAPTURE_PHOTO_DIR, file_name)
            with open(file_path, "wb") as f:
                f.write(capture_photo.file.read())
            capture_photo_path = file_name
        
        # 确定出入类型（根据时间简单判断：早上进入，晚上离开）
        current_hour = datetime.now().hour
        access_type = 1 if 6 <= current_hour < 14 else 0
        
        # 创建出入记录
        record = AccessRecord(
            staff_id=staff.id,
            device_id=device_id,
            access_type=access_type,
            access_time=datetime.now(),
            similarity=similarity,
            capture_photo_path=capture_photo_path,
            verify_result=1 if similarity >= FACE_SIMILARITY_THRESHOLD else 0
        )
        db.add(record)
        db.commit()
        db.refresh(record)
        
        return {
            "code": 200,
            "message": "验证成功" if similarity >= FACE_SIMILARITY_THRESHOLD else "验证失败",
            "data": {
                "verify_result": 1 if similarity >= FACE_SIMILARITY_THRESHOLD else 0,
                "similarity": similarity,
                "staff_id": staff.id,
                "staff_name": staff.name,
                "staff_no": staff.staff_no,
                "access_type": access_type,
                "record_id": record.id
            }
        }
    
    # 没有提供staff_id，模拟未知人员验证失败
    record = AccessRecord(
        staff_id=0,
        device_id=device_id,
        access_type=1,
        access_time=datetime.now(),
        similarity=0.0,
        verify_result=0,
        remark="未识别到人员信息"
    )
    db.add(record)
    db.commit()
    
    return {
        "code": 400,
        "message": "验证失败：未识别到人员信息",
        "data": {"verify_result": 0, "similarity": 0.0}
    }


@router.get("/export/excel")
def export_access_records(
    staff_name: Optional[str] = None,
    staff_no: Optional[str] = None,
    device_id: Optional[int] = None,
    access_type: Optional[int] = None,
    verify_result: Optional[int] = None,
    start_time: Optional[str] = None,
    end_time: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    导出出入记录到Excel
    """
    query = db.query(AccessRecord)
    
    # 应用筛选条件
    if staff_name:
        query = query.join(Staff).filter(Staff.name.contains(staff_name))
    if staff_no:
        query = query.join(Staff).filter(Staff.staff_no.contains(staff_no))
    if device_id:
        query = query.filter(AccessRecord.device_id == device_id)
    if access_type is not None:
        query = query.filter(AccessRecord.access_type == access_type)
    if verify_result is not None:
        query = query.filter(AccessRecord.verify_result == verify_result)
    if start_time:
        try:
            start_dt = datetime.strptime(start_time, "%Y-%m-%d")
            query = query.filter(AccessRecord.access_time >= start_dt)
        except ValueError:
            pass
    if end_time:
        try:
            end_dt = datetime.strptime(end_time, "%Y-%m-%d")
            end_dt = end_dt.replace(hour=23, minute=59, second=59)
            query = query.filter(AccessRecord.access_time <= end_dt)
        except ValueError:
            pass
    
    # 预加载关联对象并查询
    records = query.options(
        joinedload(AccessRecord.staff),
        joinedload(AccessRecord.device)
    ).order_by(desc(AccessRecord.access_time)).all()
    
    # 创建Excel工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "出入记录"
    
    # 定义样式
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    cell_alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入表头
    headers = ["序号", "工号", "姓名", "设备名称", "设备位置", "出入类型", "出入时间", "相似度", "验证结果", "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # 写入数据
    for row_idx, record in enumerate(records, 2):
        record_dict = record.to_dict()
        row_data = [
            row_idx - 1,
            record_dict.get("staff_no", ""),
            record_dict.get("staff_name", ""),
            record_dict.get("device_name", ""),
            record_dict.get("location", ""),
            record_dict.get("access_type_text", ""),
            record_dict.get("access_time", ""),
            record_dict.get("similarity", 0),
            record_dict.get("verify_result_text", ""),
            record_dict.get("remark", "")
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = cell_alignment
            cell.border = thin_border
    
    # 设置列宽
    column_widths = [8, 12, 12, 15, 20, 10, 20, 10, 10, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    
    # 保存到内存
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # 生成文件名
    filename = f"出入记录_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/statistics/today")
def get_today_statistics(db: Session = Depends(get_db)):
    """
    获取今日统计数据
    """
    today = date.today()
    today_start = datetime.combine(today, datetime.min.time())
    today_end = datetime.combine(today, datetime.max.time())
    
    # 今日总记录数
    total_records = db.query(AccessRecord).filter(
        AccessRecord.access_time >= today_start,
        AccessRecord.access_time <= today_end
    ).count()
    
    # 今日成功验证数
    success_records = db.query(AccessRecord).filter(
        AccessRecord.access_time >= today_start,
        AccessRecord.access_time <= today_end,
        AccessRecord.verify_result == 1
    ).count()
    
    # 今日进入人数（去重）
    enter_staff = db.query(AccessRecord.staff_id).filter(
        AccessRecord.access_time >= today_start,
        AccessRecord.access_time <= today_end,
        AccessRecord.access_type == 1,
        AccessRecord.verify_result == 1
    ).distinct().count()
    
    # 今日离开人数（去重）
    leave_staff = db.query(AccessRecord.staff_id).filter(
        AccessRecord.access_time >= today_start,
        AccessRecord.access_time <= today_end,
        AccessRecord.access_type == 0,
        AccessRecord.verify_result == 1
    ).distinct().count()
    
    return {
        "code": 200,
        "message": "success",
        "data": {
            "total_records": total_records,
            "success_records": success_records,
            "enter_count": enter_staff,
            "leave_count": leave_staff,
            "date": today.strftime("%Y-%m-%d")
        }
    }
