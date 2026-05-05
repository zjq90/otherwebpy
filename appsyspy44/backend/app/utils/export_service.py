from io import BytesIO
from typing import List, Dict
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

def create_excel_report(report_data: Dict, report_type: str, period: str) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{report_type}报表"
    
    header_font = Font(bold=True, size=14, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    
    ws.merge_cells("A1:F1")
    ws["A1"] = f"{report_type}报表 - {period}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = header_alignment
    
    ws.row_dimensions[1].height = 30
    
    if report_type == "生产数据":
        headers = ["日期", "计划产量", "实际产量", "合格产量", "完成率(%)", "合格率(%)"]
        row_data = []
        for item in report_data.get("data", []):
            row_data.append([
                item.get("date", ""),
                item.get("planned_quantity", 0),
                item.get("actual_quantity", 0),
                item.get("qualified_quantity", 0),
                item.get("completion_rate", 0),
                item.get("pass_rate", 0)
            ])
        
        summary_row = [
            "合计/平均",
            report_data.get("total_planned", 0),
            report_data.get("total_actual", 0),
            report_data.get("total_qualified", 0),
            report_data.get("avg_completion_rate", 0),
            report_data.get("avg_pass_rate", 0)
        ]
    
    elif report_type == "质量趋势":
        headers = ["日期", "总样本数", "合格样本数", "合格率(%)", "材料名称"]
        row_data = []
        for item in report_data.get("data", []):
            row_data.append([
                item.get("date", ""),
                item.get("total_samples", 0),
                item.get("passed_samples", 0),
                item.get("pass_rate", 0),
                item.get("material_name", "")
            ])
        
        summary_row = [
            "合计/平均",
            report_data.get("total_samples", 0),
            report_data.get("total_passed", 0),
            report_data.get("avg_pass_rate", 0),
            ""
        ]
    
    elif report_type == "设备开机率":
        headers = ["日期", "设备名称", "开机率(%)", "实际运行(小时)", "停机时间(小时)"]
        row_data = []
        for item in report_data.get("data", []):
            row_data.append([
                item.get("date", ""),
                item.get("equipment_name", ""),
                item.get("operating_rate", 0),
                item.get("actual_runtime", 0),
                item.get("downtime", 0)
            ])
        
        summary_row = [
            "合计/平均",
            "",
            report_data.get("avg_operating_rate", 0),
            "",
            report_data.get("total_downtime", 0)
        ]
    
    else:
        headers = ["数据"]
        row_data = []
        summary_row = [""]
    
    col_count = len(headers)
    for col in range(1, col_count + 1):
        cell = ws.cell(row=3, column=col, value=headers[col - 1])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 18
    
    for row_idx, row_item in enumerate(row_data, start=4):
        for col_idx, value in enumerate(row_item, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    summary_start_row = 4 + len(row_data)
    for col_idx, value in enumerate(summary_row, start=1):
        cell = ws.cell(row=summary_start_row, column=col_idx, value=value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output

def create_production_chart_data(report_data: Dict) -> Dict:
    return {
        "chart_type": "bar",
        "title": "产量统计图",
        "x_axis": [item.get("date", "") for item in report_data.get("data", [])],
        "series": [
            {"name": "计划产量", "data": [item.get("planned_quantity", 0) for item in report_data.get("data", [])]},
            {"name": "实际产量", "data": [item.get("actual_quantity", 0) for item in report_data.get("data", [])]},
            {"name": "合格产量", "data": [item.get("qualified_quantity", 0) for item in report_data.get("data", [])]}
        ]
    }

def create_rate_chart_data(report_data: Dict, rate_type: str = "completion") -> Dict:
    if rate_type == "completion":
        data_key = "completion_rate"
        title = "任务完成率趋势"
    else:
        data_key = "pass_rate"
        title = "生产合格率趋势"
    
    return {
        "chart_type": "line",
        "title": title,
        "x_axis": [item.get("date", "") for item in report_data.get("data", [])],
        "series": [
            {"name": "百分比(%)", "data": [item.get(data_key, 0) for item in report_data.get("data", [])]}
        ]
    }
